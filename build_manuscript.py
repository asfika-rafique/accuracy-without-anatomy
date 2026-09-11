# -*- coding: utf-8 -*-
"""Build the IEEE JBHI-format manuscript (two-column) as .docx."""
import os, json
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

import sys as _sys, os as _os
_sys.path.insert(0, _os.path.join(_os.path.dirname(_os.path.abspath(__file__)), "exp"))
from _config import ROOT, FIGS
OUT = _os.path.join(ROOT, "JBHI_Stroke_CT_REBUILT.docx")

doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Inches(8.5), Inches(11)
s.top_margin, s.bottom_margin = Inches(0.72), Inches(0.9)
s.left_margin = s.right_margin = Inches(0.625)


def set_cols(section, n, space=340):
    sectPr = section._sectPr
    cols = sectPr.find(qn('w:cols'))
    if cols is None:
        cols = OxmlElement('w:cols'); sectPr.append(cols)
    cols.set(qn('w:num'), str(n)); cols.set(qn('w:space'), str(space))
    cols.set(qn('w:equalWidth'), '1')


set_cols(s, 1)
st = doc.styles['Normal']
st.font.name = 'Times New Roman'; st.font.size = Pt(10)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
st.paragraph_format.space_before = Pt(0); st.paragraph_format.space_after = Pt(0)
st.paragraph_format.line_spacing = 1.0


def P(text='', size=10, bold=False, italic=False, align='just', indent=0.0,
      before=0, after=0, allcaps=False):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Inches(indent)
    p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(after)
    p.alignment = {'just': WD_ALIGN_PARAGRAPH.JUSTIFY, 'c': WD_ALIGN_PARAGRAPH.CENTER,
                   'l': WD_ALIGN_PARAGRAPH.LEFT}[align]
    if text:
        r = p.add_run(text); r.font.size = Pt(size); r.bold = bold; r.italic = italic
        r.font.name = 'Times New Roman'
        if allcaps: r.font.all_caps = True
    return p


def RUNS(parts, size=9, align='just', indent=0.0, after=0):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Inches(indent)
    p.paragraph_format.space_after = Pt(after)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if align == 'just' else WD_ALIGN_PARAGRAPH.CENTER
    for t, b, i in parts:
        r = p.add_run(t); r.font.size = Pt(size); r.bold = b; r.italic = i
        r.font.name = 'Times New Roman'
    return p


def H1(t): P(t, size=10, align='c', before=8, after=3, allcaps=True).paragraph_format.keep_with_next = True
def H2(t): P(t, size=10, italic=True, align='l', before=4, after=1).paragraph_format.keep_with_next = True
def BODY(t, first=False): P(t, size=10, align='just', indent=0.0 if first else 0.18)


def _newsec(n):
    sec = doc.add_section(WD_SECTION.CONTINUOUS)
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin, sec.bottom_margin = Inches(0.72), Inches(0.9)
    sec.left_margin = sec.right_margin = Inches(0.625)
    set_cols(sec, n)
    return sec


_FIGW = {"fig_audit.png": 7.16, "fig_results.png": 7.16, "fig_legacy.png": 6.4}


def FIGWIDE(fname, caption):
    _newsec(1)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3); p.paragraph_format.space_after = Pt(1)
    p.add_run().add_picture(os.path.join(FIGS, fname), width=Inches(_FIGW.get(fname, 7.16)))
    c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    c.paragraph_format.space_after = Pt(4)
    r = c.add_run(caption); r.font.size = Pt(7.5); r.font.name = 'Times New Roman'
    _newsec(2)


def _shade(cell, hexc):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear')
    sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), hexc); tcPr.append(sh)


def TABLE(num, title, rows, widths, wide=False, notes=None, fs=7.5):
    if wide: _newsec(1)
    cap = doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(5); cap.paragraph_format.space_after = Pt(1)
    cap.paragraph_format.keep_with_next = True
    r1 = cap.add_run('TABLE ' + str(num)); r1.font.size = Pt(7.5); r1.font.name = 'Times New Roman'
    cap.add_run().add_break()
    r2 = cap.add_run(title); r2.font.size = Pt(7.5); r2.font.name = 'Times New Roman'
    r2.font.all_caps = True
    t = doc.add_table(rows=0, cols=len(widths)); t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, row in enumerate(rows):
        row_obj = t.add_row()
        trpr = row_obj._tr.get_or_add_trPr()
        if ri == 0: trpr.append(OxmlElement('w:tblHeader'))
        trpr.append(OxmlElement('w:cantSplit'))
        cells = row_obj.cells
        for ci, txt in enumerate(row):
            cells[ci].width = Inches(widths[ci])
            para = cells[ci].paragraphs[0]
            para.paragraph_format.keep_with_next = ri < len(rows) - 1
            para.paragraph_format.space_before = Pt(0.5); para.paragraph_format.space_after = Pt(0.5)
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER if (ci > 0 or ri == 0) else WD_ALIGN_PARAGRAPH.LEFT
            run = para.add_run(txt); run.font.size = Pt(fs); run.font.name = 'Times New Roman'
            if ri == 0: run.bold = True; _shade(cells[ci], 'EFEFEF')
    if notes:
        n = doc.add_paragraph(); n.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        n.paragraph_format.space_after = Pt(4)
        rr = n.add_run(notes); rr.font.size = Pt(7); rr.font.name = 'Times New Roman'
    else:
        doc.add_paragraph().paragraph_format.space_after = Pt(3)
    if wide: _newsec(2)


# ============================ TITLE ==========================================
P('Accuracy Without Anatomy: A Redundancy and Shortcut Audit of a Public '
  'Head-CT Stroke Benchmark', size=17.5, align='c', after=7)
P('Tanha Asfika Jaman, Iftee Shekh Iftesham, Mst Lovely Akter, and MOSTAFA FARZANA',
  size=10.5, align='c', after=2)
P('School of Computer Science, Nanjing University of Information Science and Technology, '
  'Nanjing 210044, China', size=8.5, align='c', after=1)
P('E-mail: 202353460049@nuist.edu.cn; 202353460052@nuist.edu.cn; 202353080157@nuist.edu.cn; '
  '202353080121@nuist.edu.cn', size=8.5, align='c', after=1)
P('Corresponding author: Tanha Asfika Jaman.', size=8.5, italic=True, align='c', after=7)

s2 = doc.add_section(WD_SECTION.CONTINUOUS)
s2.page_width, s2.page_height = Inches(8.5), Inches(11)
s2.top_margin, s2.bottom_margin = Inches(0.72), Inches(0.9)
s2.left_margin = s2.right_margin = Inches(0.625)
set_cols(s2, 2)

# ============================ ABSTRACT =======================================
RUNS([('Abstract\u2014', True, True),
      ('Objective: Near-ceiling stroke classification on public head computed tomography benchmarks may not measure lesion recognition. We evaluate benchmark reliability. Methods: On TEKNOFEST-2021 (7202 slices; 1290 haemorrhagic, 1361 ischaemic, 4551 normal), we calibrated a redundancy detector on 429 same-scan pairs. Three networks used a frozen protocol, a held-out test partition, three seeds and bootstrap intervals. We compared random and redundancy-aware splits, retrained on four content streams, validated lesion containment with expert masks, and measured attribution. Results: DenseNet-201 reached 0.966 accuracy (95 per cent interval 0.960 to 0.972); ResNet-50 and ResNet-18 were within one point. Test images with a training near-duplicate scored 0.998 versus 0.969 for the rest. Deleting the geometric intracranial compartment reduced retrained accuracy from 0.961 to 0.901, although this stream retains some annotated lesion content. Exterior content yielded macro area under the curve 0.834 versus chance 0.500, with no annotated lesion pixels on 2214 scorable masks. Full-slice models scored 0.669 with the intracranial compartment deleted versus 0.566 with only that compartment retained. Permutation discrimination was near chance; duplicate exclusion preserved the main findings. Conclusion: Substantial predictive signal remains available outside the geometric intracranial compartment, including beyond the segmented head. Significance: This reproducible audit supports trustworthy medical-imaging benchmark evaluation; headline accuracy alone is not evidence of clinically meaningful lesion recognition.',
       False, False)], size=8.5, after=4)

RUNS([('Index Terms\u2014', True, True),
      ('Benchmark evaluation, brain stroke, computed tomography, data leakage, medical image '
       'analysis, shortcut learning.', False, False)], size=8.5, after=5)

# ============================ I. INTRODUCTION ================================
H1('I. Introduction')
BODY('Stroke was the third leading cause of death worldwide in 2021 and a dominant cause of '
     'acquired adult disability [1]. Non-contrast computed tomography (NCCT) is the first-line '
     'examination because it is fast, ubiquitous and sufficient to exclude haemorrhage before '
     'thrombolysis, but hyperacute ischaemic change on NCCT is a low-contrast, spatially diffuse '
     'finding.', first=True)

BODY('Deep learning has produced convincing head-CT results where the target is dense, the cohort large and the reference standard strong: detectors developed using hundreds of thousands of scans perform well on critical head-CT findings [2], and a single-centre model reached an area under the receiver operating characteristic curve (AUC) of 0.991 against neuroradiologist consensus while also localizing the bleed [3]. Other studies report stroke classification on small public NCCT collections, commonly close to ceiling — for instance 98.9 per cent accuracy for stroke detection, 98.5 per cent for ischaemia-versus-haemorrhage classification and lesion-segmentation intersection-over-union of 0.952 on Turkish Ministry of Health head-CT data [4].')

BODY('Such scores raise an evaluation question: how much reflects pathology rather than dataset structure? Medical-imaging models can exploit laterality tokens, support devices and acquisition signatures [5]–[7], while leakage between partitions inflates reported performance [8], [9]. Separating these effects is a biomedical-informatics problem because benchmarks determine which medical-AI comparisons are credible.')

BODY('We present a practical audit framework and demonstrate it on the public TEKNO21 stroke benchmark [10], rather than propose another high-accuracy classifier. The contribution combines (i) redundancy measurement calibrated on 429 same-scan pairs; (ii) leakage-aware evaluation under a frozen protocol; (iii) intracranial, complementary head-content and exterior ablations; (iv) expert-mask validation of annotated-lesion containment; and (v) multi-seed attribution, permutation and duplicate-exclusion controls. Together, these distinguish redundancy-related score inflation from predictive content outside the segmented head, primarily background and head-support content, and provide a reproducible basis for evaluating medical-imaging benchmarks. Redundancy grouping does not establish patient-level independence.')

# ============================ II. RELATED WORK ===============================
H1('II. Related Work')
BODY('Haemorrhage detection on head CT is the mature end of this field: Chilamkurthy et al. developed and validated models using 313318 scans from roughly twenty centres [2] and Kuo et al. reported expert-level performance with simultaneous localization [3].', first=True)

BODY('Shortcut learning is one explanation for high scores that fail to transfer. Geirhos et al. framed it generally [5]; Zech et al. and DeGrave et al. demonstrated shortcut signals in chest radiographs [6], [7]. Souza et al. decoded site and scanner information from a Parkinson disease classifier in JBHI [11], while Boland et al. localized shortcut representations within networks [12]. We combine content ablation with expert-mask checks to assess retained lesion content, rather than assume that a geometric deletion removes pathology.')

BODY('Leakage is a separate evaluation failure. Kapoor and Narayanan documented it across 294 papers in 17 fields [8]; imaging studies reported substantial inflation from slice-level splitting [9], [13]. More recently, Abhishek et al. audited duplication and partition leakage in dermatological benchmarks and released corrected datasets [14]. Our contribution is the joint measurement of redundancy effects and anatomy-related predictive content under one controlled stroke-CT protocol; the constituent auditing ideas are established.')

BODY('Attribution maps are common interpretability evidence [15], but across eight methods on two radiology benchmarks every method failed at least one trustworthiness criterion [16]. We therefore quantify attribution against expert masks and geometric compartments as corroborative spatial evidence, not a causal explanation.')

BODY('The TEKNOFEST-2021 (TEKNO21) collection [10] provides anonymised NCCT slices in three '
     'classes, curated from the Turkish Ministry of Health teleradiology archive and annotated '
     'by seven radiologists. It has become a common substrate for stroke classification work, '
     'and is the object of this audit.')

# ============================ III. METHODS ===================================
H1('III. Materials and Methods')

H2('A. Benchmark')
BODY('We used the public three-class TEKNO21 release [10] as distributed in Parquet form on the Hugging Face mirror named in the data statement, which preserves the original release filenames. It contains 7202 axial NCCT slices: 1290 haemorrhagic, 1361 ischaemic and 4551 normal, so the majority-class accuracy is 0.632. Images arrive as rendered 8-bit PNG, not DICOM; Hounsfield-unit windowing is therefore neither possible nor applied, although it is frequently described in papers using this release. Images were cached at 256 pixels, resized to 224 pixels, scaled to the unit interval, replicated over three channels and normalised with ImageNet statistics. Training used random resized crops (scale 0.85–1.0, aspect ratio 0.9–1.11), horizontal flips and rotations up to 7 degrees. The release carries no subject identifiers.', first=True)

H2('B. Redundancy Audit')
BODY('429 images are named "{class}_dcm_{ID}.png" and every one of those IDs also occurs as '
     '"{ID}.png" with an identical class label: the same scan is present twice under two '
     'filenames, rendered through two paths. These 429 pairs are ground truth for calibrating a '
     'near-duplicate detector.', first=True)

BODY('We computed a 256-bit difference hash per image. The 429 known pairs have median Hamming '
     'distance 0 (maximum 5); random pairs have median 102, with a first percentile of 66 '
     '[Fig. 1(a), (b)]. The two distributions are completely separated, so any threshold in a '
     'wide range is defensible; we used 6, the smallest value recovering 100 per cent of the '
     'known pairs [Fig. 1(c)]. Connected components of the resulting graph define redundancy '
     'groups.')

BODY('We also tested whether the original file numbering encodes examination structure, which would imply same-patient slices. Mean cosine similarity between ImageNet-embedded images with adjacent identifiers is 0.6236, against a random-pair baseline of 0.6235. This does not support filename adjacency as a proxy for study membership and does not establish patient independence. Because subject identifiers are absent, these groups remain a lower bound on true examination-level grouping.')

FIGWIDE('fig_audit.png',
 'Fig. 1.  Redundancy audit and content streams. (a) The 429 scans that the release stores twice under two filenames give ground-truth same-scan pairs; their Hamming distances (median 0) are completely separated from random pairs (median 102), so the near-duplicate detector is calibrated rather than assumed. (b) One such pair. (c) Redundancy and known-pair recall against threshold; we use 6, the smallest value recovering all known pairs, at which 15.5 per cent of the collection lies in a redundancy group. (d)–(g) The four content streams: full slice; intracranial compartment only; the head with that compartment deleted; and the region outside the segmented head; lesion exclusion is assessed against expert masks.')



H2('C. Split Policies and Protocol')
BODY('Two policies were compared, identical in every other respect — architecture, epochs, batch size, learning rate, schedule and seeds, verified from the saved run records. The random policy performs an image-level stratified 70/15/15 split, which is what image-level evaluation on this collection amounts to. The grouped policy assigns whole redundancy groups, keeping every detected redundancy group inside one partition. Under the random policy 15.6 to 16.1 per cent of test images have a training near-duplicate; under the grouped policy this is exactly zero, with class balance preserved to within one image per class.', first=True)

BODY('The full protocol is given in Table I and was applied uniformly in the final runs: class-weighted cross-entropy, AdamW [17], cosine schedule, mixed precision [18], model selection on validation macro-F1 only, and the test partition scored once with the selected epoch. Every configuration was run with three seeds; the split is regenerated per seed, so the reported spread includes split variability. Interval estimates use two bootstrap procedures: 5000 resamples from pooled seed predictions for the single-configuration intervals of Table II, and 3333 paired resamples within each seed, combined into 9999 draws, for the contrasts of Table III, where pairing removes between-image variance from the difference. Training used an NVIDIA RTX 4060 Ti (8 GB) with PyTorch 2.5.1; cuDNN autotuning and mixed precision make runs reproducible only to within the seed spread, not bitwise.')

TABLE('I', 'Benchmark Composition and Frozen Evaluation Protocol',
[["Item", "Value"],
 ["Collection", "TEKNO21 public three-class NCCT release [10]"],
 ["Images", "7202 axial slices, 8-bit PNG"],
 ["Classes (h / i / n)", "1290 / 1361 / 4551"],
 ["Majority-class accuracy", "0.632"],
 ["Provenance duplicates", "429 scans stored twice under two filenames"],
 ["Near-duplicate detector", "256-bit dHash, threshold 6"],
 ["   known-pair recall at threshold", "100 %  (median distance 0; random pairs 102)"],
 ["Images in a redundancy group", "1118  (15.5 %)"],
 ["Split", "70 / 15 / 15, stratified, three seeds"],
 ["   random policy: test leakage", "15.6\u201316.1 % of test has a train near-duplicate"],
 ['   grouped: detected-duplicate leakage', "0 %"],
 ["Model selection", "validation macro-F1 only"],
 ["Test partition", "scored once, with the selected epoch"],
 ["Optimiser / schedule", 'AdamW [17], cosine, mixed precision [18]'],
 ["Loss", "class-weighted cross-entropy"],
 ["Epochs / batch size / learning rate", "15 / 48 / 3e-4, identical for every run"],
 ["Hardware / framework", "RTX 4060 Ti 8 GB, PyTorch 2.5.1+cu121"],
 ["Intervals", '5000-sample bootstrap (Table II); 9999-draw within-seed paired bootstrap (Table III)']],
 [1.55, 1.75], wide=False, fs=7,
 notes='h = haemorrhagic, i = ischaemic, n = normal. The protocol was applied uniformly in the final runs and is identical across every configuration reported.')



H2('D. Content Ablation')
BODY('We built four input streams from each slice [Fig. 1(d)–(g)]. Full slice is unmodified. Intracranial-only retains the geometric region enclosed by the bright skull ring; it is not a brain-tissue segmentation. For compactness, figures and tables label it brain or brain-only. Non-brain denotes the complementary head content after deleting that compartment, including skull, scalp, orbits and skull base; it may retain lesion pixels. Exterior deletes the segmented head, retaining primarily background and head-support content; it is distinct from the non-brain stream. Expert masks quantify annotated-lesion containment (Section IV-C). Networks were retrained from ImageNet initialisation on each stream under the grouped policy. DenseNet-201 [19], ResNet-50 and ResNet-18 [20], spanning 11.2 to 23.5 million parameters, served as capacity controls. Reporting follows CLAIM [21]; result provenance is recorded below.', first=True)

BODY('The claim that a stream excludes the lesion is geometric, so we validated it against expert '
     'annotation. A separately hosted mirror of the same collection distributes stroke '
     'segmentation masks; 6643 of its 6650 numeric filenames match ours, confirming a common '
     'source. Its mask semantics check out: all 1093 haemorrhagic and all 1130 ischaemic masks '
     'are non-empty, and all 4427 normal masks are empty. We resampled the 2223 lesion masks to '
     'our grid, of which 2214 survived a minimum-area threshold, and measured what fraction of '
     'annotated lesion pixels falls inside each of our masks (Section IV-C).')

BODY('Two controls accompany the ablation. Permuting all labels tests whether the pipeline can '
     'manufacture apparent performance. Repeating the ablation with the 429 duplicated images '
     'removed tests whether any effect is an artefact of that differently rendered subset.')

H2('E. Attribution')
BODY('Grad-CAM [15] for the predicted class was computed on the last convolutional block of all three selected full-slice ResNet-18 checkpoints. Outside-intracranial attribution used 1061, 1069 and 1069 test slices for seeds 0, 1 and 2, respectively, with intracranial masks of at least 100 pixels and nonzero CAM mass. Lesion attribution used 335, 329 and 333 annotated test slices with masks of at least 20 pixels. We report mean and sample SD across seeds, comparing attribution mass with the corresponding image-area fraction.', first=True)

# ============================ IV. RESULTS ====================================
H1('IV. Results')

H2('A. Reproducible Near-Ceiling Classification')
BODY('Under the redundancy-aware protocol with a held-out test partition, DenseNet-201 reached 0.9661 ± 0.0053 accuracy (bootstrap interval 0.9596 to 0.9722), macro-F1 0.9551 and macro-AUC 0.9953 (Table II). ResNet-50 reached 0.9695 and ResNet-18 0.9605; the spread across three architectures whose parameter counts differ by a factor of 2.1 (11.2, 18.1 and 23.5 million) is 0.89 points, comparable to the within-architecture seed spread [Fig. 2(c)]. Near-ceiling performance is reproduced across these three tested backbones; this does not establish equivalence or generalize to other architectures.',
     first=True)

FIGWIDE('fig_results.png',
 'Fig. 2.  (a) Content ablation under the redundancy-aware protocol (ResNet-18, three seeds, error bars are SD). Deleting the intracranial compartment costs only 5.9 accuracy points, and the region outside the head alone still gives macro-AUC 0.834 against 0.500 for chance. (b) Redundancy leakage on the random split: test images with a training near-duplicate are classified far more accurately than the rest, although this subset is small enough that the aggregate split effect is 0.8 to 1.3 points. (c) Three architectures spanning 11.2 to 23.5 million parameters agree to within 0.89 points. Within these tested backbones, architecture is not a dominant determinant of performance.')

TABLE('II', 'Test-Set Results Under the Redundancy-Aware Protocol',
[["Configuration", "Input", "Accuracy", "Macro-F1", "Macro-AUC"],
 ["DenseNet-201", "full slice", "0.9661 \u00b1 0.0053", "0.9551 \u00b1 0.0081", "0.9953 \u00b1 0.0015"],
 ["ResNet-50", "full slice", "0.9695 \u00b1 0.0009", "0.9615 \u00b1 0.0019", "0.9937 \u00b1 0.0020"],
 ["ResNet-18", "full slice", "0.9605 \u00b1 0.0095", "0.9473 \u00b1 0.0130", "0.9919 \u00b1 0.0031"],
 ["ResNet-18", "brain only", "0.9056 \u00b1 0.0028", "0.8773 \u00b1 0.0068", "0.9693 \u00b1 0.0069"],
 ["ResNet-18", "non-brain (brain compartment deleted)", "0.9010 \u00b1 0.0088", "0.8773 \u00b1 0.0097", "0.9681 \u00b1 0.0065"],
 ["ResNet-18", "exterior (head mask deleted)", "0.6616 \u00b1 0.0273", "0.5985 \u00b1 0.0249", "0.8342 \u00b1 0.0148"],
 ["control: permuted labels", "full slice", "0.4476 \u00b1 0.0240", "0.3427 \u00b1 0.0131", "0.5107 \u00b1 0.0131"],
 ["control: duplicates removed", "full slice", "0.9616 \u00b1 0.0020", "0.9453 \u00b1 0.0053", "0.9940 \u00b1 0.0008"],
 ["control: duplicates removed", "non-brain", "0.9043 \u00b1 0.0134", "0.8720 \u00b1 0.0195", "0.9694 \u00b1 0.0071"],
 ["control: duplicates removed", "exterior", "0.6512 \u00b1 0.0267", "0.5816 \u00b1 0.0252", "0.8223 \u00b1 0.0191"],
 ["reference: majority class", "\u2013", "0.6319", "0.2582", "0.5000"],
 ["Random-split comparison", "", "", "", ""],
 ["DenseNet-201, random split", "full slice", "0.9739 \u00b1 0.0046", "0.9659 \u00b1 0.0052", "0.9966 \u00b1 0.0020"],
 ["   with train near-duplicate", "full slice", "0.9980", "\u2013", "\u2013"],
 ["   without near-duplicate", "full slice", "0.9693", "\u2013", "\u2013"],
 ["ResNet-18, random split", "full slice", "0.9739 \u00b1 0.0014", "0.9666 \u00b1 0.0019", "0.9947 \u00b1 0.0043"],
 ["   with train near-duplicate", "full slice", "0.9961", "\u2013", "\u2013"],
 ["   without near-duplicate", "full slice", "0.9697", "\u2013", "\u2013"]],
 [1.75, 1.55, 1.30, 1.25, 1.25], wide=True, fs=7,
 notes='Mean \u00b1 SD over three seeds (the split is regenerated per seed, so the spread includes '
       'split variability). Bootstrap 95 % intervals for the full-slice grouped runs: DenseNet-201 '
       'accuracy 0.9596\u20130.9722; ResNet-50 0.9633\u20130.9753; ResNet-18 0.9534\u20130.9670. The '
       'stratified rows partition the random-split test set by whether an image has a near-duplicate '
       'in training (n \u2248 171 and 912 respectively).')



H2('B. Redundancy Inflates Accuracy, Most Sharply on the Affected Subset')
BODY('Comparing split policies, the aggregate effect is present but modest: 0.9739 random against 0.9661 '
     'grouped for DenseNet-201 (+0.78 points) and 0.9739 against 0.9605 for ResNet-18 '
     '(+1.33 points). Both architectures are inflated by a redundancy-blind split, by between three quarters of a point and one and a third.', first=True)

BODY('Stratifying within the random split is more informative [Fig. 2(b)]. Test images that have a near-duplicate in training are classified at 0.9980 (DenseNet-201) and 0.9961 (ResNet-18), against 0.9693 and 0.9697 for the rest: gaps of 2.87 and 2.64 points on the affected subset. The aggregate difference is smaller; the affected subset is about a sixth of the test set and performance is already near ceiling. The redundancy is real and its local effect is measurable; on a harder task, or a benchmark with more duplication, the aggregate difference could be larger.')

H2('C. Accuracy Persists After Intracranial Deletion')
BODY('Deleting anatomy and retraining gives the central result [Fig. 2(a), Table II]. The full slice yields 0.9605 accuracy and macro-AUC 0.9919. Restricting the input to the geometric intracranial mask lowers accuracy to 0.9056. Deleting that geometric compartment and retaining complementary head content lowers it only to 0.9010 with macro-AUC 0.9681 — within half a point of the intracranial-only stream, and still 0.90 accuracy on a task whose majority class is 0.632.',
     first=True)

BODY('Validation against expert masks changes how much weight the non-brain stream can carry, and '
     'we report it plainly. Across 2214 annotated slices, a mean of 62.8 per cent of lesion '
     'pixels fall inside our geometric intracranial mask (median 92.2 per cent), but only '
     '51.0 per cent of slices reach 90 per cent containment, and the lowest decile is near '
     'zero \u2014 the skull-base slices where the bone ring is open. The non-brain stream is '
     'therefore not a clean lesion-free input: for a substantial minority of slices it retains '
     'part of the lesion, and its 0.9010 accuracy must be read with that caveat.')

BODY('The exterior stream is not affected by this, and it is the decisive one. Every annotated lesion pixel in all 2214 slices lies inside the head mask — containment is 1.000 — so a model restricted to the exterior sees no annotated lesion pixels on the 2214 scorable slices. With the segmented head removed, the three-way label is still recovered at macro-AUC 0.8342 ± 0.0148 and accuracy 0.6616. Chance macro-AUC is 0.5. Content containing no annotated lesion pixels in the scorable masks therefore carries substantial label information. Acquisition, cropping and rendering are plausible sources; the physical source remains unidentified.')

BODY('The controls support this interpretation. Permuting labels collapses the same pipeline, across three seeds, to 0.4476 ± 0.0240 accuracy and macro-AUC 0.5107 ± 0.0131, providing no evidence of above-chance discrimination under label permutation. Removing the 429 duplicated images preserves the main pattern: full 0.9616, non-brain 0.9043 (macro-AUC 0.9694), exterior 0.6512 (macro-AUC 0.8223). The exterior signal is therefore not confined to the duplicated subset.')


BODY('These contrasts used paired resampling within each seed, with the resulting draws combined across seeds (Table III). Restricting the input to the intracranial compartment costs 5.48 accuracy points (95 per cent interval 3.43 to 7.68, p < 0.001). Deleting that compartment instead costs only 0.47 points relative to keeping it and nothing else (interval −2.04 to +2.96, p = 0.75). The interval excludes any advantage for the intracranial-only stream larger than about three points, so this is a bounded null rather than a bare absence of significance. The exterior stream sits 23.9 points below the non-brain stream, but its macro-AUC interval of 0.801 to 0.862 excludes chance by a wide margin. Table III reports four contrasts; the three significant ones remain significant under Bonferroni correction for that family, and the fourth is a null that no correction can affect. The Table III intervals are percentiles of the combined within-seed bootstrap draws, reflecting test-image resampling and differences among the three fitted runs; they are not confidence intervals for a seed-averaged or patient-level effect. Zero observed tail counts are reported conservatively as p < 0.001. Table II separately reports seed SDs.')

BODY('Because the intracranial mask degrades where the skull ring is open, we repeated the comparison on the 986 test slices per seed whose intracranial mask exceeds 2 per cent of the image, a stricter criterion than the 100-pixel floor used for attribution. The ordering is unchanged: full 0.9605, intracranial-only 0.9233, non-brain 0.8954, exterior 0.6510. Excluding unreliable slices helps the intracranial-only stream, as expected, and opens a modest 2.8-point gap over the non-brain stream. Under either analysis the conclusion is the same: high accuracy persists after geometric intracranial deletion, subject to residual lesion content.')

TABLE('III', 'Paired Bootstrap Contrasts on Test Accuracy',
[["Contrast", "Difference", "95 % interval", "p"],
 ["full slice \u2212 brain only", "+0.0548", "+0.0343 to +0.0768", '< 0.001'],
 ["brain only \u2212 non-brain", "+0.0047", "\u22120.0204 to +0.0296", "0.75"],
 ["full slice \u2212 non-brain", "+0.0597", "+0.0306 to +0.0851", '< 0.001'],
 ["non-brain \u2212 exterior", "+0.2392", "+0.1739 to +0.2923", '< 0.001'],
 ["exterior macro-AUC vs chance", "0.8342", "0.8007 to 0.8620", "chance = 0.500"]],
 [1.55, 0.85, 1.25, 0.85], wide=False, fs=7,
 notes='3333 paired resamples within each seed, combined into 9999 draws across three seeds (ResNet-18, redundancy-aware split). The second row is the key result: deleting the segmented intracranial compartment does not significantly change accuracy relative to keeping only that compartment.')

BODY('Two further analyses sharpen the interpretation. First, the per-class breakdown (Table IV) shows the effect is not confined to the subtle class. For haemorrhage, the class with the most conspicuous imaging finding, the non-brain stream reaches recall 0.866 and AUC 0.972, marginally above the intracranial-only stream (0.854, 0.972): geometric intracranial deletion preserves haemorrhage classification performance despite possible residual lesion content. Ischaemia behaves the same way (0.827 versus 0.825). Only the full-slice model, which sees both, does materially better.')

BODY('Full-slice and non-brain models agree on 89.0 per cent of test images and are both correct on 87.8 per cent; the full-slice model alone is correct on 8.2 per cent. Agreement is compatible with shared predictive information, but cannot establish which features either model uses.')

BODY('Image geometry alone offers a limited explanation: 98.2, 97.1 and 95.9 per cent of haemorrhagic, ischaemic and normal images respectively are 512 by 512, so this coarse frame-size feature does not explain near-ceiling discrimination. The content ablations provide the stronger evidence for label information outside the brain.')

TABLE('IV', 'Per-Class Content Ablation (ResNet-18, Grouped Split, Three Seeds Pooled)',
[["Input stream", "Haemorrhagic", "Ischaemic", "Normal"],
 ["full slice", "0.942 / 0.995", "0.887 / 0.988", "0.988 / 0.992"],
 ["brain only", "0.854 / 0.972", "0.825 / 0.965", "0.944 / 0.971"],
 ["non-brain (brain compartment deleted)", "0.866 / 0.972", "0.827 / 0.966", "0.933 / 0.965"],
 ["exterior (head mask deleted)", "0.505 / 0.813", "0.697 / 0.838", "0.695 / 0.850"]],
 [1.60, 1.05, 1.05, 1.05], wide=False, fs=7,
 notes='Each cell is recall / one-versus-rest AUC. For haemorrhage — the class with the most conspicuous imaging finding — deleting the intracranial compartment does not reduce performance below the intracranial-only stream. Chance AUC is 0.500.')




H2('D. Full-Slice Models on Ablated Inputs')
BODY('The ablation above retrains on each stream, so it measures how much information a stream contains. A complementary question is whether the trained full-slice model actually uses that information. We therefore took the three trained full-slice models and, without any retraining, evaluated them on the ablated inputs.', first=True)

BODY('The ordering is clear. On complete slices these models score 0.9599; given only the intracranial compartment they fall to 0.5660 (macro-AUC 0.7185), and given the head with the intracranial compartment deleted they reach 0.6690 (macro-AUC 0.7793). Accuracy is higher with the intracranial compartment deleted than with that compartment alone. With the segmented head deleted it falls to 0.3356, below the majority-class rate, showing that the exterior-trained signal does not transfer directly to the full-slice model. All masked conditions are out-of-distribution for a model trained on complete slices, so the absolute values are not comparable with the retrained streams of Section IV-C; their ordering shows greater retained accuracy for the non-brain input but does not quantify causal feature importance.')

H2('E. Attribution Relative to Lesions and Intracranial Area')
BODY('With expert masks available we can score attribution against the lesion itself rather than against the intracranial mask, and we repeat it on all three full-slice checkpoints rather than one. Grad-CAM places 3.83 ± 0.22 per cent of its mass inside the annotated lesion, against a lesion area share of 1.43 per cent: an enrichment of 2.68 ± 0.16 times over a uniform null. Attribution is therefore lesion-directed to a measurable but modest degree, and 96 per cent of its mass still falls outside the lesion.')

BODY('Measured against the intracranial mask instead, Grad-CAM placed 66.6 ± 0.6 per cent of its mass outside the intracranial compartment, where uniform attribution would place 84.8 per cent: a measured-to-uniform ratio of 0.79. Attribution is thus biased toward the intracranial compartment relative to its area, although two thirds of its mass lies outside; this spatial distribution is consistent with the ablation findings but does not establish causal dependence. Both quantities are stable across the three seeds. Table V summarizes result provenance.', first=True)

TABLE('V', 'Provenance of Every Reported Quantity',
[["Result", "Category", "Basis"],
 ["Benchmark composition, class counts, duplicate structure", "new measurement", "computed from the public release"],
 ["Near-duplicate calibration and grouping", "new measurement", "429 known same-scan pairs as ground truth"],
 ["Split-policy comparison and stratified leakage", "new controlled experiment", "3 seeds, frozen protocol, held-out test"],
 ["Architecture comparison", "new controlled experiment", "3 seeds each, bootstrap intervals"],
 ["Content ablation and both controls", "new controlled experiment", "3 seeds per stream, ImageNet initialization"],
 ["Per-class ablation, model agreement, geometry check", "new measurement", "derived from saved test predictions"],
 ["Lesion containment and attribution-vs-lesion", "new measurement", '2214 scorable masks; attribution on 335/329/333 test slices'],
 ["Cross-stream dependence of the trained model", "new controlled experiment", "3 checkpoints, inference only"],
 ['Grad-CAM mass outside intracranial mask', "new measurement", '1061/1069/1069 test slices, three selected checkpoints']],
 [2.55, 1.45, 2.60], wide=True, fs=7,
 notes='Every experimental quantity in this paper was produced in this study. A “new measurement” is computed directly from the released data or from stored test predictions; a “new controlled experiment” includes training or inference under the frozen protocol of Table I.')



# ============================ V. DISCUSSION ==================================
H1('V. Discussion')
BODY('The demonstrated finding is that high TEKNO21 classification scores coexist with substantial prediction from content outside the geometric intracranial compartment. The tested backbones achieve near-ceiling full-slice accuracy; deleting that compartment retains 0.901 accuracy, with the residual-lesion caveat. Exterior macro-AUC remains 0.834 despite no annotated lesion pixels on the 2214 scorable masks. Its 0.6616 accuracy is only modestly above the 0.6319 majority-class reference; discrimination and accuracy should not be conflated.', first=True)

BODY('Cross-stream inference provides complementary evidence: full-slice models retain more accuracy with the intracranial compartment deleted than with it alone. These masked inputs are out of distribution; their ordering does not identify causal features or establish that lesion information is unused.')

BODY('A plausible explanation is an acquisition- or presentation-level signature. Scanner, protocol, cropping or head-support differences could correlate with class and supply predictive cues, consistent with shortcut findings in other imaging tasks [6], [7], [11]. These factors were not measured here. The audit establishes retained predictive signal, but scanner metadata and controlled acquisition comparisons would be needed to identify its physical source.')

BODY('The biomedical-informatics implication concerns the validity of model comparisons. High scores can reflect dataset structure rather than pathology, so architecture rankings alone cannot establish progress in lesion recognition. Within the tested backbones, the 0.89-point spread provides limited evidence that architecture is a dominant performance determinant. A content ablation with independently checked lesion containment makes the benchmark more informative by measuring prediction that survives deletion of the targeted content.')

BODY('Our redundancy finding is more nuanced than the leakage literature might predict. The '
     'duplication is unambiguous \u2014 429 scans stored twice, and 15.5 per cent of images in a '
     'near-duplicate group \u2014 and its effect on the affected subset is large, at about three '
     'points. The aggregate effect is smaller, 0.8 to 1.3 points, because that subset is about a sixth of the test set and the task saturates. Reporting '
     'only the aggregate difference between split policies would have hidden a real defect; the '
     'stratified comparison is the more sensitive instrument and we would encourage its use.')

BODY('The audit suggests three practical checks for similar medical-image benchmarks: calibrate redundancy detection and report performance separately for test images with training near-duplicates; retrain on defined content streams and validate annotated-lesion containment; and quantify attribution relative to lesion and compartment area across seeds. Reporting these alongside aggregate scores helps distinguish dataset redundancy, residual lesion content and exterior predictive signal. These checks complement patient-disjoint and external evaluation; they do not replace them.')

BODY('This is a biomedical-AI evaluation study, not a clinical deployment study. Stroke imaging can influence acute management decisions, but no reader comparison, clinical calibration, prospective evaluation or external validation was performed. The contribution concerns reliable medical-imaging benchmark design and interpretation: headline accuracy is not, by itself, evidence of clinically meaningful lesion recognition.')

# ============================ VI. LIMITATIONS ================================
H1('VI. Limitations')
BODY('The intracranial masks are geometric, not expert tissue segmentations, and degrade at the skull base; 384 essentially empty masks were retained in training but excluded from attribution. Expert-mask validation shows residual lesion pixels in the non-brain stream for a substantial minority of slices. Restricting analysis to reliable masks preserves the main ordering (Section IV-C). The exterior stream contains no annotated lesion pixels on 2214 scorable masks, but this does not prove absence of all patient tissue or unannotated lesions.', first=True)

BODY('Subject-level separation cannot be verified, since the release omits identifiers. Our grouping is a calibrated lower bound on it, and the embedding analysis does not identify study membership, but a true patient-level split remains unavailable and would require reliable patient identifiers.')

BODY('The exterior ablation establishes predictive information outside the segmented head; it does not identify the physical factor supplying it or prove absence of all patient tissue. Scanner and protocol metadata needed to test acquisition explanations are unavailable.')

BODY("Findings are confined to one collection from one national health system and to 2D slice classification; neither the effect sizes nor the audit's performance have been externally validated. The reproducible workflow can be adapted to similar benchmarks, but its masks and redundancy calibration require dataset-specific validation. Three seeds provide a limited sample of training/split variability, and runs are not bitwise reproducible.")

# ============================ VII. CONCLUSION ================================
H1('VII. Conclusion')
BODY('We demonstrated a reproducible redundancy and content-audit framework on TEKNO21. Near-ceiling accuracy persists across the three tested backbones; deleting the geometric intracranial compartment retains 0.901 accuracy, while exterior content reaches macro-AUC 0.834 with no annotated lesion pixels in the 2214 scorable masks. Random versus grouped splitting increases accuracy by 0.8 to 1.3 points overall; training-near-duplicate test images score about three points higher than the remainder. These findings show why benchmark scores alone do not establish lesion recognition. Calibrated redundancy checks, validated content ablations and multi-seed controls offer a practical basis for more trustworthy medical-imaging evaluation.', first=True)

# ============================ STATEMENTS =====================================
H1('Data, Code and Ethics Statements')
BODY('Data: the public TEKNO21 head-CT stroke collection [10], obtained from the Hugging Face mirror BTX24/tekno21-brain-stroke-dataset-multi, which carries an Apache-2.0 repository tag; the underlying collection was released by the Turkish Ministry of Health open-data portal and its conditions of use are set by the issuing institutions rather than by the mirror, so intending users should confirm them at source. The expert stroke masks used only for the containment and attribution validation of Section IV-C come from a second mirror of the same collection, Karrar-Alhdrawi/brain-stroke-ct-dataset, matched to ours by original filename. Only illustrative CT panels are included; no source image dataset is redistributed and no new patient data were collected. Ethics: analysis used only anonymised, publicly released images; ethical approval for the original collection and annotation is reported in [10]. Code: all audit, training, ablation and analysis scripts, together with the two tools used to verify protocol uniformity across runs and to cross-check every number reported here against the stored result files, are available at https://github.com/asfika-rafique/accuracy-without-anatomy. The authors declare no competing interests.', first=True)

# ============================ REFERENCES =====================================
H1('References')
refs = ['GBD 2021 Stroke Risk Factor Collaborators, "Global, regional, and national burden of stroke and its risk factors, 1990–2021: a systematic analysis for the Global Burden of Disease Study 2021," Lancet Neurol., vol. 23, no. 10, pp. 973–1003, 2024.', 'S. Chilamkurthy et al., "Deep learning algorithms for detection of critical findings in head CT scans: a retrospective study," Lancet, vol. 392, no. 10162, pp. 2388–2396, 2018.', 'W. Kuo, C. Häne, P. Mukherjee, J. Malik, and E. L. Yuh, "Expert-level detection of acute intracranial hemorrhage on head computed tomography using deep learning," Proc. Natl. Acad. Sci. USA, vol. 116, no. 45, pp. 22737–22745, 2019.', 'S. Yalçın and H. Vural, "Brain stroke classification and segmentation using encoder-decoder based deep convolutional neural networks," Comput. Biol. Med., vol. 149, art. 105941, 2022.', 'R. Geirhos et al., "Shortcut learning in deep neural networks," Nature Mach. Intell., vol. 2, pp. 665–673, 2020.', 'J. R. Zech, M. A. Badgeley, M. Liu, A. B. Costa, J. J. Titano, and E. K. Oermann, "Variable generalization performance of a deep learning model to detect pneumonia in chest radiographs: A cross-sectional study," PLoS Med., vol. 15, no. 11, art. e1002683, 2018.', 'A. J. DeGrave, J. D. Janizek, and S.-I. Lee, "AI for radiographic COVID-19 detection selects shortcuts over signal," Nature Mach. Intell., vol. 3, pp. 610–619, 2021.', 'S. Kapoor and A. Narayanan, "Leakage and the reproducibility crisis in machine-learning-based science," Patterns, vol. 4, no. 9, art. 100804, 2023.', 'E. Yagis et al., "Effect of data leakage in brain MRI classification using 2D convolutional neural networks," Sci. Rep., vol. 11, art. 22544, 2021.', 'U. Koç et al., "Artificial intelligence in healthcare competition (TEKNOFEST-2021): Stroke data set," Eurasian J. Med., vol. 54, no. 3, pp. 248–258, 2022.', 'R. Souza et al., "Identifying biases in a multicenter MRI database for Parkinson\'s disease classification: Is the disease classifier a secret site classifier?," IEEE J. Biomed. Health Inform., vol. 28, no. 4, pp. 2047–2054, 2024, doi: 10.1109/JBHI.2024.3352513.', 'C. Boland, K. A. Goatman, S. A. Tsaftaris, and S. Dahdouh, "There are no shortcuts to anywhere worth going: Identifying shortcuts in deep learning models for medical image analysis," in Proc. MIDL, PMLR, vol. 250, pp. 131–150, 2024.', 'I. E. Tampu, A. Eklund, and N. Haj-Hosseini, "Inflation of test accuracy due to data leakage in deep learning-based classification of OCT images," Sci. Data, vol. 9, art. 580, 2022.', 'K. Abhishek, A. Jain, and G. Hamarneh, "Investigating the quality of DermaMNIST and Fitzpatrick17k dermatological image datasets," Sci. Data, vol. 12, art. 196, 2025, doi: 10.1038/s41597-025-04382-5.', 'R. R. Selvaraju, M. Cogswell, A. Das, R. Vedantam, D. Parikh, and D. Batra, "Grad-CAM: Visual explanations from deep networks via gradient-based localization," in Proc. IEEE ICCV, 2017, pp. 618–626.', 'N. Arun et al., "Assessing the trustworthiness of saliency maps for localizing abnormalities in medical imaging," Radiol. Artif. Intell., vol. 3, no. 6, art. e200267, 2021.', 'I. Loshchilov and F. Hutter, "Decoupled weight decay regularization," in Proc. ICLR, 2019.', 'P. Micikevicius et al., "Mixed precision training," in Proc. ICLR, 2018.', 'G. Huang, Z. Liu, L. van der Maaten, and K. Q. Weinberger, "Densely connected convolutional networks," in Proc. IEEE CVPR, 2017, pp. 4700–4708.', 'K. He, X. Zhang, S. Ren, and J. Sun, "Deep residual learning for image recognition," in Proc. IEEE CVPR, 2016, pp. 770–778.', 'J. Mongan, L. Moy, and C. E. Kahn, "Checklist for artificial intelligence in medical imaging (CLAIM): A guide for authors and reviewers," Radiol. Artif. Intell., vol. 2, no. 2, art. e200029, 2020.']
for i, r in enumerate(refs, 1):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Inches(-0.15)
    p.paragraph_format.left_indent = Inches(0.15)
    p.paragraph_format.space_after = Pt(0.5)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run('[%d] %s' % (i, r))
    run.font.size = Pt(7.5); run.font.name = 'Times New Roman'

# A trailing continuous section break makes Word balance the two columns of the
# reference list instead of leaving the right-hand column of the last page empty.
_bal = doc.add_section(WD_SECTION.CONTINUOUS)
_bal.page_width, _bal.page_height = Inches(8.5), Inches(11)
_bal.top_margin, _bal.bottom_margin = Inches(0.72), Inches(0.9)
_bal.left_margin = _bal.right_margin = Inches(0.625)
set_cols(_bal, 2)

url = 'https://github.com/asfika-rafique/accuracy-without-anatomy'
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from copy import deepcopy
for paragraph in doc.paragraphs:
    for run in list(paragraph.runs):
        if url in run.text:
            before, after = run.text.split(url, 1)
            run.text = before
            h = OxmlElement('w:hyperlink')
            h.set(qn('r:id'), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
            rr = OxmlElement('w:r')
            if run._r.rPr is not None: rr.append(deepcopy(run._r.rPr))
            tt = OxmlElement('w:t'); tt.text = url; rr.append(tt); h.append(rr)
            run._r.addnext(h)
            tail = deepcopy(run._r)
            for child in list(tail):
                if child.tag != qn('w:rPr'): tail.remove(child)
            tt = OxmlElement('w:t'); tt.set(qn('xml:space'), 'preserve'); tt.text = after; tail.append(tt)
            h.addnext(tail)
doc.save(OUT)
print('saved', OUT)
